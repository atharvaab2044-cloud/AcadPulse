"""Private, expiring preview storage and safe CSV/XLSX parsing for bulk workflows.

Previews are held as random-token JSON files under UPLOAD_ROOT, never under static.
They are bound to the initiating user and target record and expire after two hours.
The committing route must revalidate every row against current database state.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import secrets
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from werkzeug.utils import secure_filename

MAX_ROWS = 5000
MAX_BYTES = 16 * 1024 * 1024
PREVIEW_TTL_HOURS = 2
TOKEN_RE = re.compile(r"^[a-f0-9]{40}$")
HEADER_ALIASES = {
    "student_id": "student_code", "student_code": "student_code", "student_code_id": "student_code",
    "roll": "roll_number", "roll_no": "roll_number", "roll_number": "roll_number", "rollno": "roll_number",
    "prn": "prn", "permanent_registration_number": "prn", "permanent_registration_no": "prn",
    "marks": "marks", "mark": "marks", "marks_obtained": "marks", "score": "marks",
    "status": "status", "attendance": "status", "attendance_status": "status", "mark_status": "status",
    "correction_reason": "correction_reason", "reason": "correction_reason", "note": "correction_reason",
}


class SpreadsheetImportError(ValueError):
    pass


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value):
    raw = _cell_text(value).lstrip("\ufeff").strip().lower()
    normalized = re.sub(r"[^a-z0-9]+", "_", raw).strip("_")
    return HEADER_ALIASES.get(normalized, normalized)


def parse_spreadsheet(content: bytes, filename: str, *, workflow: str):
    """Parse first worksheet/CSV into normalized records, never matching students here."""
    if not content:
        raise SpreadsheetImportError("The uploaded spreadsheet is empty.")
    if len(content) > MAX_BYTES:
        raise SpreadsheetImportError("The spreadsheet exceeds the 16 MB import limit.")
    suffix = Path(secure_filename(filename or "")).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise SpreadsheetImportError("Upload a .csv or .xlsx file. Legacy .xls and macro-enabled files are not accepted.")
    if suffix == ".csv":
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise SpreadsheetImportError("CSV must be UTF-8 encoded. Save it as CSV UTF-8 and retry.") from exc
        reader = csv.reader(io.StringIO(decoded, newline=""))
        try:
            raw_headers = next(reader)
        except StopIteration as exc:
            raise SpreadsheetImportError("The CSV has no header row.") from exc
        raw_rows = [(line_number, row) for line_number, row in enumerate(reader, start=2) if any(str(c or "").strip() for c in row)]
    else:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise SpreadsheetImportError("The XLSX file could not be opened. It may be corrupt or encrypted.") from exc
        try:
            sheet = workbook[workbook.sheetnames[0]]
            iterator = sheet.iter_rows(values_only=True)
            try:
                raw_headers = list(next(iterator))
            except StopIteration as exc:
                raise SpreadsheetImportError("The first worksheet is empty.") from exc
            raw_rows = []
            for line_number, row in enumerate(iterator, start=2):
                values = list(row)
                if any(value is not None and str(value).strip() for value in values):
                    raw_rows.append((line_number, values))
                if len(raw_rows) > MAX_ROWS:
                    raise SpreadsheetImportError(f"A maximum of {MAX_ROWS} data rows can be imported at once.")
        finally:
            workbook.close()
    if len(raw_rows) > MAX_ROWS:
        raise SpreadsheetImportError(f"A maximum of {MAX_ROWS} data rows can be imported at once.")

    headers = [_header(value) for value in raw_headers]
    if not headers or not any(headers):
        raise SpreadsheetImportError("A header row is required.")
    if len(headers) != len(set(headers)):
        raise SpreadsheetImportError("The spreadsheet contains duplicate/ambiguous column headers after normalization.")
    required = {"status"}
    if workflow == "coursework":
        required.add("marks")
    if not {"student_code", "roll_number", "prn"}.intersection(headers):
        raise SpreadsheetImportError("Include at least one identity column: student_code, roll_number, or verified PRN.")
    missing = required - set(headers)
    if missing:
        raise SpreadsheetImportError("Missing required column(s): " + ", ".join(sorted(missing)))
    records = []
    for line_number, values in raw_rows:
        if len(values) > len(headers) and any(v is not None and str(v).strip() for v in values[len(headers):]):
            record = {key: "" for key in headers}
            records.append({"row_number": line_number, "values": record,
                            "parse_error": "This row has more values than the header row."})
            continue
        padded = values + [None] * (len(headers) - len(values))
        normalized = {key: _cell_text(padded[idx]) for idx, key in enumerate(headers)}
        records.append({"row_number": line_number, "values": normalized})
    if not records:
        raise SpreadsheetImportError("The spreadsheet has a header but no data rows.")
    return records


def store_import_preview(*, upload_root, workflow, target_id, actor_id, filename, content, rows):
    root = Path(upload_root).resolve()
    staging_dir = root / "import_previews"
    source_dir = root / f"{workflow}_imports"
    staging_dir.mkdir(parents=True, exist_ok=True)
    source_dir.mkdir(parents=True, exist_ok=True)
    token = secrets.token_hex(20)
    digest = hashlib.sha256(content).hexdigest()
    safe_name = secure_filename(filename) or f"import{Path(filename or '').suffix.lower()}"
    suffix = Path(safe_name).suffix.lower()
    source_key = f"{workflow}_imports/{secrets.token_hex(20)}{suffix}"
    source_path = (root / source_key).resolve()
    if root not in source_path.parents:
        raise SpreadsheetImportError("Invalid upload storage path.")
    source_path.write_bytes(content)
    now = datetime.now(timezone.utc)
    payload = {
        "token": token, "workflow": workflow, "target_id": int(target_id), "actor_id": int(actor_id),
        "source_filename": safe_name[:255], "source_storage_key": source_key, "source_sha256": digest,
        "created_at": now.isoformat(), "expires_at": (now + timedelta(hours=PREVIEW_TTL_HOURS)).isoformat(),
        "rows": rows,
    }
    preview_path = staging_dir / f"{token}.json"
    preview_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return token, payload


def read_import_preview(*, upload_root, token, workflow, target_id, actor_id):
    if not token or not TOKEN_RE.fullmatch(token):
        raise SpreadsheetImportError("The import preview token is invalid. Upload the spreadsheet again.")
    root = Path(upload_root).resolve()
    path = (root / "import_previews" / f"{token}.json").resolve()
    if root not in path.parents or not path.is_file():
        raise SpreadsheetImportError("The import preview has expired or was not found. Upload the spreadsheet again.")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        expiry = datetime.fromisoformat(payload["expires_at"])
    except (OSError, ValueError, KeyError, TypeError) as exc:
        path.unlink(missing_ok=True)
        raise SpreadsheetImportError("The import preview is invalid. Upload the spreadsheet again.") from exc
    if expiry.tzinfo is None:
        expiry = expiry.replace(tzinfo=timezone.utc)
    if datetime.now(timezone.utc) >= expiry:
        path.unlink(missing_ok=True)
        raise SpreadsheetImportError("The import preview expired after two hours. Upload the spreadsheet again.")
    if payload.get("workflow") != workflow or payload.get("target_id") != int(target_id) or payload.get("actor_id") != int(actor_id):
        raise SpreadsheetImportError("This preview belongs to another user or target record.")
    return path, payload


def delete_import_preview(path):
    Path(path).unlink(missing_ok=True)
