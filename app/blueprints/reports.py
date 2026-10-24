"""Descriptive analytics and CSV/XLSX/PDF exports with transparent labels."""
import csv
import io
from collections import defaultdict
from decimal import Decimal

from flask import Blueprint, abort, make_response, send_file
from flask_login import login_required
from sqlalchemy import select, func
from sqlalchemy.orm import selectinload

from ..extensions import db
from ..models import (
    Assessment, AssessmentMark, AttendanceRecord, AttendanceSession, ExamAttempt,
    Student, StudentEnrollment, StudentIdentifier,
)

bp = Blueprint("reports", __name__)


def report_rows():
    enrollments = db.session.execute(select(StudentEnrollment).options(
        selectinload(StudentEnrollment.student).selectinload(Student.identifiers),
        selectinload(StudentEnrollment.program), selectinload(StudentEnrollment.curriculum_version)
    ).order_by(StudentEnrollment.academic_year, StudentEnrollment.semester_number, StudentEnrollment.division,
               StudentEnrollment.roll_number, StudentEnrollment.student_id)).scalars().all()
    attendance_counts = defaultdict(lambda: {"present": 0, "absent": 0, "excused": 0})
    attendance_data = db.session.execute(select(AttendanceRecord.enrollment_id, AttendanceRecord.status).join(
        AttendanceSession, AttendanceSession.id == AttendanceRecord.session_id)).all()
    for enrollment_id, status in attendance_data:
        attendance_counts[enrollment_id][status] += 1
    marks_data = db.session.execute(select(AssessmentMark.enrollment_id, AssessmentMark.status,
        AssessmentMark.marks_obtained, Assessment.max_marks).join(Assessment, Assessment.id == AssessmentMark.assessment_id)).all()
    marks_counts = defaultdict(lambda: {"sum_pct": Decimal("0"), "count": 0})
    for enrollment_id, status, marks, maximum in marks_data:
        if status == "present" and marks is not None and maximum:
            marks_counts[enrollment_id]["sum_pct"] += (Decimal(marks) / Decimal(maximum)) * Decimal("100")
            marks_counts[enrollment_id]["count"] += 1
    attempts_data = db.session.execute(select(ExamAttempt.enrollment_id, ExamAttempt.student_id, ExamAttempt.verification_status,
        ExamAttempt.official_status, ExamAttempt.attempt_type)).all()
    attempts_counts = defaultdict(lambda: {"verified": 0, "verified_fail_attempts": 0})
    # Match attempt to the exact enrollment when available; students with historical attempts
    # lacking enrollment IDs are intentionally not assigned to an arbitrary term.
    for enrollment_id, student_id, verification_status, official_status, attempt_type in attempts_data:
        if enrollment_id and verification_status == "verified":
            attempts_counts[enrollment_id]["verified"] += 1
            if official_status == "fail":
                attempts_counts[enrollment_id]["verified_fail_attempts"] += 1
    output = []
    for enrollment in enrollments:
        prn = next((i.identifier_value for i in enrollment.student.identifiers
                    if i.identifier_type == "prn" and i.verification_status == "verified"), "")
        a = attendance_counts[enrollment.id]
        attendance_total = a["present"] + a["absent"]
        attendance_percent = round(100 * a["present"] / attendance_total, 2) if attendance_total else None
        m = marks_counts[enrollment.id]
        coursework_mean = round(float(m["sum_pct"] / m["count"]), 2) if m["count"] else None
        e = attempts_counts[enrollment.id]
        output.append({
            "student_code": enrollment.student.student_code,
            "student_name": enrollment.student.full_name,
            "verified_prn": prn,
            "programme_code": enrollment.program.code,
            "programme_name": enrollment.program.name,
            "academic_year": enrollment.academic_year,
            "semester": enrollment.semester_number,
            "roll_number": enrollment.roll_number or "",
            "division": enrollment.division or "",
            "practical_batch": enrollment.practical_batch or "",
            "enrollment_type": enrollment.enrollment_type,
            "enrollment_status": enrollment.enrollment_status,
            "curriculum_version": enrollment.curriculum_version.version_code if enrollment.curriculum_version else "",
            "attendance_percent": attendance_percent,
            "coursework_mean_percent": coursework_mean,
            "verified_exam_attempts": e["verified"],
            "verified_fail_attempts_historical": e["verified_fail_attempts"],
        })
    return output


COLUMNS = [
    ("student_code", "Student code"), ("student_name", "Student name"), ("verified_prn", "Verified PRN"),
    ("programme_code", "Programme code"), ("programme_name", "Programme name"), ("academic_year", "Academic year"),
    ("semester", "Semester"), ("roll_number", "Roll number"), ("division", "Division"),
    ("practical_batch", "Practical batch"), ("enrollment_type", "Enrollment type"),
    ("enrollment_status", "Enrollment status"), ("curriculum_version", "Curriculum version"),
    ("attendance_percent", "Attendance % (recorded present/absent sessions; excused excluded)"),
    ("coursework_mean_percent", "Coursework mean % (descriptive)"),
    ("verified_exam_attempts", "Verified exam attempts"),
    ("verified_fail_attempts_historical", "Verified failed attempts (historical; not an official backlog calculation)"),
]


def _text_safe(value):
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


@bp.get("/analytics/export.csv")
@login_required
def export_csv():
    rows = report_rows()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([label for _, label in COLUMNS])
    for row in rows:
        writer.writerow([_text_safe(row[key]) for key, _ in COLUMNS])
    response = make_response("\ufeff" + output.getvalue())
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = 'attachment; filename="acadpulse-academic-overview.csv"'
    return response


@bp.get("/analytics/export.xlsx")
@login_required
def export_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        abort(501, "Excel export requires openpyxl. Install requirements-document-processing.txt.")
    rows = report_rows()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Academic overview"
    sheet.append([label for _, label in COLUMNS])
    for row in rows:
        sheet.append([_text_safe(row[key]) for key, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3D68")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 40
    for col_idx, (key, label) in enumerate(COLUMNS, start=1):
        values = [str(sheet.cell(row=r, column=col_idx).value or "") for r in range(1, min(sheet.max_row, 200) + 1)]
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max((len(v) for v in values), default=len(label)) + 2, 12), 40)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="acadpulse-academic-overview.xlsx")


@bp.get("/analytics/export.pdf")
@login_required
def export_pdf():
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    except ImportError:
        abort(501, "PDF export requires ReportLab. Install requirements-document-processing.txt.")
    rows = report_rows()
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A3), rightMargin=8*mm, leftMargin=8*mm, topMargin=8*mm, bottomMargin=8*mm)
    styles = getSampleStyleSheet()
    story = [Paragraph("AcadPulse — Academic Overview", styles["Title"]),
             Paragraph("Descriptive export generated from stored enrollment, attendance, coursework and verified-attempt records. It is not an official promotion/backlog decision.", styles["BodyText"]), Spacer(1, 4*mm)]
    data = [[label for _, label in COLUMNS]]
    for row in rows:
        data.append(["—" if row[key] is None else str(row[key]) for key, _ in COLUMNS])
    width_mm = [30, 38, 26, 24, 42, 24, 14, 18, 14, 20, 22, 22, 32, 29, 28, 20, 34]
    width_scale = min(1.0, 404 / sum(width_mm))
    widths = [w * width_scale * mm for w in width_mm]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3D68")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6), ("LEADING", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FB")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(table)
    doc.build(story)
    output.seek(0)
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name="acadpulse-academic-overview.pdf")
